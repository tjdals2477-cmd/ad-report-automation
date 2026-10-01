from __future__ import annotations

import csv
import io
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook


MAX_ROWS_PER_DATASET = 100_000


ALIASES = {
    "campaign": ["캠페인명", "캠페인", "campaignname", "campaign"],
    "spend": [
        "총광고비",
        "광고비",
        "광고비용",
        "집행금액",
        "소진액",
        "비용",
        "adspend",
        "spend",
        "cost",
    ],
    "sales": [
        "총전환매출액14일",
        "총전환매출14일",
        "광고매출액",
        "광고매출",
        "전환매출액",
        "전환매출",
        "총매출",
        "매출액",
        "매출",
        "adsales",
        "revenue",
        "sales",
    ],
    "clicks": ["클릭수", "클릭", "유입수", "clicks", "click"],
    "orders": [
        "총주문수14일",
        "구매수",
        "주문수",
        "전환수",
        "총주문수",
        "구매건수",
        "conversions",
        "orders",
        "purchases",
    ],
    "conversion_rate": [
        "전환율",
        "구매전환율",
        "cvr",
        "conversionrate",
    ],
    "option_id": [
        "옵션id",
        "옵션아이디",
        "노출상품id",
        "판매상품id",
        "vendoritemid",
        "optionid",
        "sku",
    ],
    "product_name": [
        "상품명",
        "노출상품명",
        "광고상품명",
        "판매상품명",
        "제품명",
        "productname",
        "product",
        "itemname",
    ],
    "quantity": [
        "총판매수량14일",
        "판매수량",
        "주문수량",
        "구매수량",
        "총판매수량",
        "수량",
        "units",
        "quantity",
        "qty",
    ],
}


@dataclass
class Dataset:
    dataset_id: str
    file_name: str
    sheet_name: str
    columns: list[str]
    rows: list[dict[str, Any]]

    def public_summary(self) -> dict[str, Any]:
        return {
            "id": self.dataset_id,
            "label": f"{self.file_name} · {self.sheet_name}",
            "fileName": self.file_name,
            "sheetName": self.sheet_name,
            "columns": self.columns,
            "rowCount": len(self.rows),
            "preview": self.rows[:5],
            "detected": detect_columns(self.columns),
        }


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", str(value or "").strip().lower())


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"열 {index}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    return headers


def _find_header_row(matrix: list[list[Any]]) -> int:
    best_index = 0
    best_score = -1
    for index, row in enumerate(matrix[:20]):
        non_empty = sum(value not in (None, "") for value in row)
        text_like = sum(bool(isinstance(value, str) and value.strip()) for value in row)
        score = non_empty * 2 + text_like
        if non_empty >= 2 and score > best_score:
            best_index = index
            best_score = score
    return best_index


def _matrix_to_dataset(
    matrix: list[list[Any]], dataset_id: str, file_name: str, sheet_name: str
) -> Dataset | None:
    if not matrix:
        return None
    header_index = _find_header_row(matrix)
    headers = _unique_headers(matrix[header_index])
    records: list[dict[str, Any]] = []
    for values in matrix[header_index + 1 : header_index + 1 + MAX_ROWS_PER_DATASET]:
        padded = list(values) + [None] * (len(headers) - len(values))
        record = {headers[i]: padded[i] for i in range(len(headers))}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    if not records:
        return None
    return Dataset(dataset_id, file_name, sheet_name, headers, records)


def read_xlsx(data: bytes, file_name: str, id_prefix: str) -> list[Dataset]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    datasets: list[Dataset] = []
    try:
        for index, sheet in enumerate(workbook.worksheets):
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            dataset = _matrix_to_dataset(
                matrix,
                f"{id_prefix}:{index}",
                file_name,
                sheet.title,
            )
            if dataset:
                datasets.append(dataset)
    finally:
        workbook.close()
    return datasets


def read_delimited(data: bytes, file_name: str, id_prefix: str) -> list[Dataset]:
    text = None
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("파일 문자 인코딩을 읽을 수 없습니다.")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if "\t" in sample else ","
    matrix = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    dataset = _matrix_to_dataset(matrix, f"{id_prefix}:0", file_name, "데이터")
    return [dataset] if dataset else []


def read_upload(data: bytes, file_name: str, id_prefix: str) -> list[Dataset]:
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension in {"xlsx", "xlsm"}:
        return read_xlsx(data, file_name, id_prefix)
    if extension in {"csv", "tsv", "txt"}:
        return read_delimited(data, file_name, id_prefix)
    if extension == "xls":
        raise ValueError("구형 .xls 파일은 Excel에서 .xlsx로 저장한 뒤 올려주세요.")
    raise ValueError("지원 형식은 .xlsx, .xlsm, .csv, .tsv입니다.")


def detect_columns(columns: list[str]) -> dict[str, str | None]:
    normalized = {column: _normalize_header(column) for column in columns}
    result: dict[str, str | None] = {}
    for field, aliases in ALIASES.items():
        normalized_aliases = [_normalize_header(alias) for alias in aliases]
        # Alias order is meaningful. Specific 14-day metrics must win over
        # generic partial matches such as "전환매출" or "총주문수".
        exact = next(
            (
                column
                for alias in normalized_aliases
                for column, value in normalized.items()
                if value == alias
            ),
            None,
        )
        if exact:
            result[field] = exact
            continue
        partial = next(
            (
                column
                for alias in normalized_aliases
                for column, value in normalized.items()
                if alias and alias in value
            ),
            None,
        )
        result[field] = partial
    return result


def find_column(columns: list[str], aliases: list[str]) -> str | None:
    """Find a column by ordered aliases without falling back to other metrics."""
    normalized = {column: _normalize_header(column) for column in columns}
    normalized_aliases = [_normalize_header(alias) for alias in aliases]
    for alias in normalized_aliases:
        for column, value in normalized.items():
            if value == alias:
                return column
    for alias in normalized_aliases:
        for column, value in normalized.items():
            if alias and alias in value:
                return column
    return None


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return 0.0 if not math.isfinite(float(value)) else float(value)
    text = str(value).strip()
    if not text or text in {"-", "--", "N/A", "n/a", "없음"}:
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text.replace(",", ""))
    try:
        number = float(cleaned)
    except ValueError:
        return 0.0
    return -abs(number) if negative else number


def _percent(value: Any) -> float | None:
    if value is None or value == "":
        return None
    is_percent_text = isinstance(value, str) and "%" in value
    number = _number(value)
    if not is_percent_text and -1 <= number <= 1:
        number *= 100
    return number


def _text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    return text if text else fallback


def _mapped_value(row: dict[str, Any], mapping: dict[str, str | None], field: str) -> Any:
    column = mapping.get(field)
    return row.get(column) if column else None


def _conversion_rate(
    orders: float,
    clicks: float,
    rate_weighted_sum: float,
    rate_weight: float,
    rate_sum: float,
    rate_count: int,
) -> float | None:
    if clicks > 0 and orders >= 0:
        return orders / clicks * 100
    if rate_weight > 0:
        return rate_weighted_sum / rate_weight
    if rate_count:
        return rate_sum / rate_count
    return None


def build_report(
    campaign_dataset: Dataset,
    product_dataset: Dataset,
    campaign_mapping: dict[str, str | None],
    product_mapping: dict[str, str | None],
    meta: dict[str, Any],
) -> dict[str, Any]:
    for required in ("campaign", "spend", "sales"):
        if not campaign_mapping.get(required):
            raise ValueError(f"캠페인 데이터의 '{required}' 열을 지정해주세요.")
    for required in ("product_name", "sales"):
        if not product_mapping.get(required):
            raise ValueError(f"상품 데이터의 '{required}' 열을 지정해주세요.")

    campaign_groups: dict[str, dict[str, float]] = defaultdict(
        lambda: {
            "spend": 0.0,
            "sales": 0.0,
            "clicks": 0.0,
            "orders": 0.0,
            "rateWeightedSum": 0.0,
            "rateWeight": 0.0,
            "rateSum": 0.0,
            "rateCount": 0.0,
        }
    )
    for row in campaign_dataset.rows:
        name = _text(_mapped_value(row, campaign_mapping, "campaign"), "캠페인 미지정")
        group = campaign_groups[name]
        spend = _number(_mapped_value(row, campaign_mapping, "spend"))
        sales = _number(_mapped_value(row, campaign_mapping, "sales"))
        clicks = _number(_mapped_value(row, campaign_mapping, "clicks"))
        orders = _number(_mapped_value(row, campaign_mapping, "orders"))
        rate = _percent(_mapped_value(row, campaign_mapping, "conversion_rate"))
        group["spend"] += spend
        group["sales"] += sales
        group["clicks"] += clicks
        group["orders"] += orders
        if rate is not None:
            group["rateSum"] += rate
            group["rateCount"] += 1
            if clicks > 0:
                group["rateWeightedSum"] += rate * clicks
                group["rateWeight"] += clicks

    total_spend = sum(group["spend"] for group in campaign_groups.values())
    ad_conversion_sales = sum(group["sales"] for group in campaign_groups.values())
    # The legacy local form has no total-sales input. Keep its existing value
    # until that form is updated; the Streamlit form passes the key explicitly.
    if "totalSales" in meta:
        total_sales_input = meta["totalSales"]
        total_sales = round(float(total_sales_input)) if total_sales_input is not None else None
    else:
        total_sales = round(ad_conversion_sales)
    total_clicks = sum(group["clicks"] for group in campaign_groups.values())
    total_orders = sum(group["orders"] for group in campaign_groups.values())
    total_rate_weighted = sum(group["rateWeightedSum"] for group in campaign_groups.values())
    total_rate_weight = sum(group["rateWeight"] for group in campaign_groups.values())
    total_rate_sum = sum(group["rateSum"] for group in campaign_groups.values())
    total_rate_count = int(sum(group["rateCount"] for group in campaign_groups.values()))

    campaigns = []
    for name, group in campaign_groups.items():
        campaigns.append(
            {
                "campaign": name,
                "spend": round(group["spend"]),
                "sales": round(group["sales"]),
                "conversionRate": _conversion_rate(
                    group["orders"],
                    group["clicks"],
                    group["rateWeightedSum"],
                    group["rateWeight"],
                    group["rateSum"],
                    int(group["rateCount"]),
                ),
                "spendShare": (group["spend"] / total_spend * 100) if total_spend else 0,
            }
        )
    campaigns.sort(key=lambda item: (-item["sales"], item["campaign"]))

    product_groups: dict[tuple[str, str], dict[str, Any]] = defaultdict(
        lambda: {"sales": 0.0, "quantity": 0.0, "campaignSales": defaultdict(float)}
    )
    for row in product_dataset.rows:
        product_name = _text(_mapped_value(row, product_mapping, "product_name"))
        if not product_name:
            continue
        raw_option_id = _mapped_value(row, product_mapping, "option_id")
        option_id = (
            str(int(raw_option_id))
            if isinstance(raw_option_id, float) and raw_option_id.is_integer()
            else _text(raw_option_id, "-")
        )
        key = (option_id, product_name)
        row_sales = _number(_mapped_value(row, product_mapping, "sales"))
        product_groups[key]["sales"] += row_sales
        product_groups[key]["quantity"] += _number(
            _mapped_value(row, product_mapping, "quantity")
        )
        campaign_name = _text(_mapped_value(row, product_mapping, "campaign"))
        if campaign_name:
            product_groups[key]["campaignSales"][campaign_name] += row_sales

    products = [
        {
            "optionId": option_id,
            "campaignNames": [
                name
                for name, sales in sorted(
                    values["campaignSales"].items(), key=lambda item: (-item[1], item[0])
                )
                if sales > 0
            ] or sorted(values["campaignSales"]),
            "productName": product_name,
            "sales": round(values["sales"]),
            "quantity": round(values["quantity"]),
        }
        for (option_id, product_name), values in product_groups.items()
    ]
    products.sort(key=lambda item: (-item["sales"], item["productName"]))

    conversion_rate = _conversion_rate(
        total_orders,
        total_clicks,
        total_rate_weighted,
        total_rate_weight,
        total_rate_sum,
        total_rate_count,
    )
    source_names = sorted({campaign_dataset.file_name, product_dataset.file_name})
    return {
        "meta": {
            "periodType": meta.get("periodType", "monthly"),
            "periodLabel": _text(meta.get("periodLabel"), "기간 미지정"),
            "reportTitle": _text(meta.get("reportTitle"), "광고 성과 리포트"),
            "sourceNames": source_names,
        },
        "kpis": {
            "totalSpend": round(total_spend),
            "adConversionSales": round(ad_conversion_sales),
            "totalSales": total_sales,
            "conversionRate": conversion_rate,
            "roas": (ad_conversion_sales / total_spend * 100) if total_spend else None,
            "spendToSales": (total_spend / ad_conversion_sales * 100) if ad_conversion_sales else None,
        },
        "campaigns": campaigns,
        "products": products[:10],
        "quality": {
            "campaignRows": len(campaign_dataset.rows),
            "productRows": len(product_dataset.rows),
            "conversionMethod": (
                "주문수 ÷ 클릭수"
                if campaign_mapping.get("orders") and campaign_mapping.get("clicks")
                else "업로드 파일의 전환율"
                if campaign_mapping.get("conversion_rate")
                else "전환율 데이터 없음"
            ),
        },
    }
